from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont
import sys
P=Presentation("Score-Machine-Masterclass.pptx")
SW,SH=13.333,7.5; PX=1400; SC=PX/SW; PY_=int(SH*SC)
def inch(v): return v/914400
def font(sz,bold):
    for f in (["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"] if bold else ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]):
        try: return ImageFont.truetype(f,max(8,int(sz*SC/72)))
        except: pass
    return ImageFont.load_default()
def rgb(c):
    try: return tuple(int(c[i:i+2],16) for i in (0,2,4))
    except: return (0,0,0)
want=[int(x) for x in sys.argv[1].split(",")]
for idx in want:
    sl=P.slides[idx-1]
    master='dark' if 'DARK' in (sl.slide_layout.name or '')+(sl.slide_layout.slide_master.name or '') else 'light'
    try: bg=Image.open(f"bg_{master}.jpg").resize((PX,PY_))
    except: bg=Image.new("RGB",(PX,PY_),(255,255,255))
    img=bg.copy(); d=ImageDraw.Draw(img,"RGBA")
    for sh in sl.shapes:
        if sh.left is None: continue
        x,y,w,h=[v*SC for v in (inch(sh.left),inch(sh.top),inch(sh.width),inch(sh.height))]
        st=str(sh.shape_type)
        if sh.has_text_frame and sh.text_frame.text.strip():
            fill=None
            try:
                if sh.fill.type is not None and sh.fill.type==1: fill=sh.fill.fore_color.rgb
            except: pass
            if fill: d.rounded_rectangle([x,y,x+w,y+h],8,fill=rgb(str(fill)))
            para=sh.text_frame.paragraphs[0]
            r=para.runs[0] if para.runs else None
            sz=(r.font.size.pt if r and r.font.size else 18); bold=bool(r and r.font.bold)
            col=rgb(str(r.font.color.rgb)) if (r and r.font.color and r.font.color.type is not None) else (20,20,20)
            F=font(sz,bold); ty=y+4
            for line in sh.text_frame.text.split("\n"):
                d.text((x+4,ty),line,font=F,fill=col); ty+=F.size*1.15
            d.rectangle([x,y,x+w,y+h],outline=(255,0,255,90))
        else:
            try:
                if sh.fill.type==1: d.rounded_rectangle([x,y,x+w,y+h],8,fill=rgb(str(sh.fill.fore_color.rgb)))
            except: pass
    img.save(f"pv-{idx:02d}.png")
    print(f"pv-{idx:02d}.png")
